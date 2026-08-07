from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_PATH = (
    Path(__file__).resolve().parents[1]
    / "reports"
    / "docs"
    / "Attribution_Agent_Documentation.docx"
)

BRAND_PURPLE = RGBColor(0x78, 0x60, 0xFC)
PAPER = "F7F6F3"
STONE = "E4E4E4"
INK = RGBColor(0x1C, 0x1C, 0x1C)
MUTED = RGBColor(0x58, 0x58, 0x58)
LIGHT_PURPLE = "EEEAFD"
WHITE = "FFFFFF"

TABLE_WIDTH = 9360


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color: str = "D9D9D9", size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths: list[int], indent: int = 120) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    existing_grid = tbl.find(qn("w:tblGrid"))
    if existing_grid is not None:
        tbl.remove(existing_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl.insert(1, grid)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_border(cell)


def set_run_font(run, *, size: float | None = None, color: RGBColor | None = None, bold=None, italic=None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_format(paragraph, *, before=0, after=6, line_spacing=1.25) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line_spacing


def add_para(doc: Document, text: str = "", *, style: str | None = None, size: float | None = 11, bold=False, italic=False, color: RGBColor | None = INK, after=6, before=0, align=None):
    p = doc.add_paragraph(style=style)
    set_paragraph_format(p, before=before, after=after)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_format(p, after=4)
    run = p.add_run(text)
    set_run_font(run, size=10.5, color=INK)


def add_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    set_paragraph_format(p, after=4)
    run = p.add_run(text)
    set_run_font(run, size=10.5, color=INK)


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading("", level=level)
    if level == 1:
        set_paragraph_format(p, before=18, after=10)
        size = 16
        color = BRAND_PURPLE
    elif level == 2:
        set_paragraph_format(p, before=14, after=7)
        size = 13
        color = BRAND_PURPLE
    else:
        set_paragraph_format(p, before=10, after=5)
        size = 12
        color = RGBColor(0x38, 0x32, 0x7B)
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=True)
    return p


def add_callout(doc: Document, title: str, text: str, *, fill: str = LIGHT_PURPLE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_table_geometry(table, [TABLE_WIDTH], indent=120)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    set_paragraph_format(p, after=2)
    r = p.add_run(title)
    set_run_font(r, size=10.5, bold=True, color=INK)
    p2 = cell.add_paragraph()
    set_paragraph_format(p2, after=0)
    r2 = p2.add_run(text)
    set_run_font(r2, size=10.5, color=INK)
    add_para(doc, "", after=3)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_table_geometry(table, widths, indent=120)

    for idx, text in enumerate(headers):
        cell = table.cell(0, idx)
        set_cell_shading(cell, LIGHT_PURPLE)
        p = cell.paragraphs[0]
        set_paragraph_format(p, after=0)
        run = p.add_run(text)
        set_run_font(run, size=9.5, color=INK, bold=True)

    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            set_cell_margins(cell)
            set_cell_border(cell)
            p = cell.paragraphs[0]
            set_paragraph_format(p, after=0)
            run = p.add_run(value)
            set_run_font(run, size=9.3, color=INK)
    set_table_geometry(table, widths, indent=120)
    add_para(doc, "", after=4)
    return table


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def set_running_footer(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_format(footer, after=0)
        run = footer.add_run("ARIE Build Documentation | N8iV Promotions | Internal Operator Reference")
        set_run_font(run, size=8.5, color=MUTED)


def add_cover(doc: Document) -> None:
    add_para(doc, "N8iV Promotions", size=11, bold=True, color=BRAND_PURPLE, after=18)
    add_para(doc, "ARIE Build Documentation", size=28, bold=True, color=INK, after=4)
    add_para(
        doc,
        "Automatic Revenue Intelligence Engine / Attribution Agent",
        size=14,
        color=MUTED,
        after=18,
    )
    add_para(
        doc,
        "A compact operator guide documenting what is built, where it lives, how the pipeline runs, what is production-ready, and what still needs attention before paid pilots scale.",
        size=11.5,
        color=INK,
        after=18,
    )
    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Prepared for", "N8iV Promotions"],
            ["Build name", "ARIE - Automatic Revenue Intelligence Engine"],
            ["Workspace", r"C:\Users\zajen\attribution-agent"],
            ["Git baseline", "main at 332e165, Productize ARIE command center (#66)"],
            ["Snapshot date", datetime.now().strftime("%B %d, %Y")],
            ["Confidentiality", "Internal operator documentation. No credentials, tokens, or client secrets are included."],
        ],
        [1900, 7460],
    )
    add_callout(
        doc,
        "Executive readout",
        "ARIE is no longer just a prototype. The repository contains a managed attribution service architecture with a Streamlit Command Center, Cloud Run execution path, Databricks persistence, source connectors, attribution models, client onboarding, reporting, API access, and Telegram operator control. The remaining work is mostly verification, credential hygiene, operational hardening, and pilot onboarding.",
    )
    doc.add_page_break()


def add_document_map(doc: Document) -> None:
    add_heading(doc, "Document Map", 1)
    sections = [
        "1. Current Build Snapshot",
        "2. Product Definition",
        "3. System Architecture",
        "4. Attribution Pipeline",
        "5. Data Layer And Tables",
        "6. Command Center And Client Setup",
        "7. API, A2A, MCP, And Telegram Interfaces",
        "8. Integrations And Credentials",
        "9. Deployment And Operations",
        "10. Tests, Usability, And Known Issues",
        "11. Production Readiness Checklist",
        "12. Useful Commands",
    ]
    for section in sections:
        add_bullet(doc, section)
    add_para(
        doc,
        "Style note: this is a compact reference guide using N8iV brand accent colors. It is meant to be useful while operating the system, not just attractive as a static report.",
        size=10.5,
        color=MUTED,
        after=12,
    )


def section_current_snapshot(doc: Document) -> None:
    add_heading(doc, "1. Current Build Snapshot", 1)
    add_callout(
        doc,
        "Build status",
        "The merged base is on main at origin/main. Local work adds Telegram duplicate-send protection, Telegram event logging, desktop launcher scripts, and usability testing support. Local app services were not reachable during this documentation pass, so existing PID files should be treated as stale until ARIE is restarted.",
    )
    add_table(
        doc,
        ["Area", "Current state", "Operator meaning"],
        [
            ["Git baseline", "main...origin/main at 332e165, Productize ARIE command center (#66).", "The major ARIE Command Center productization work has been merged."],
            ["Local uncommitted changes", "Modified app/bot/Databricks writer/start-stop scripts plus new docs and launcher/test scripts.", "These are important finishing changes, especially for Telegram and desktop persistence."],
            ["Runtime check", "Streamlit on 127.0.0.1:8501 and FastAPI on 127.0.0.1:8081 were not reachable.", "Run arie_start.ps1 or the desktop shortcut before testing the UI locally."],
            ["GCP tooling", "GCP is authenticated in the operator PowerShell terminal. The Codex sandbox did not have gcloud on PATH during this documentation pass.", "Run gcloud/preflight commands from your normal PowerShell terminal; treat sandbox PATH issues as environment-specific."],
            ["Usability smoke test", "Latest recorded run: 38 pass, 2 warn, 0 fail.", "Core UI surfaces loaded on desktop and mobile; live pipeline/email actions were intentionally not triggered."],
            ["Public N8iV website", "The live URL previously showed an old/placeholder site copy path.", "DNS/Vercel/Wix transfer should be verified separately from ARIE backend readiness."],
        ],
        [1600, 3860, 3900],
    )
    add_heading(doc, "Tracked Local Work", 2)
    add_table(
        doc,
        ["File or folder", "Purpose"],
        [
            [r"attribution_agent\attribution_agent\agents\control\arie_bot.py", "Adds durable Telegram event logging, cross-process listener lock, startup events, stale update offset handling, and duplicate-response prevention."],
            [r"attribution_agent\attribution_agent\app.py", "Adds Command Center controls and Observability support for recent Telegram triggers; avoids auto-starting the bot unless explicitly configured."],
            [r"attribution_agent\attribution_agent\utils\databricks_writer.py", "Adds telegram_events table support plus write/fetch helpers."],
            [r"arie_start.ps1 / arie_stop.ps1", "Improves local process start/stop handling, PID files, hidden windows, and PATH normalization."],
            [r"scripts\arie_open_command_center.ps1", "Starts/opens ARIE in local, Cloud Run proxy, or Cloud URL mode."],
            [r"scripts\install_arie_desktop_shortcut.ps1", "Creates a persistent desktop shortcut for ARIE Command Center."],
            [r"scripts\arie_usability_smoke_test.mjs", "Playwright smoke test for command center usability across desktop and mobile."],
        ],
        [3000, 6360],
    )


def section_product(doc: Document) -> None:
    add_heading(doc, "2. Product Definition", 1)
    add_para(
        doc,
        "ARIE, the Automatic Revenue Intelligence Engine, is a managed B2B attribution service. Its job is to connect ad engagement and spend to closed-won revenue so an operator can tell a business or agency which channels, campaigns, and journeys are creating pipeline and revenue instead of only reporting clicks, leads, or platform-reported conversions.",
    )
    add_table(
        doc,
        ["Dimension", "Definition"],
        [
            ["Primary customer", "B2B businesses, marketing agencies, and agency sub-accounts that need closed-revenue attribution."],
            ["Primary operator", "N8iV as the managed-service operator from the internal Command Center."],
            ["Primary source of truth", "HubSpot closed-won deals and lifecycle events."],
            ["Revenue enrichment", "Stripe collected revenue, refunds, and true ROI when available."],
            ["Ad platforms in v1", "Meta, Google Ads, and LinkedIn Ads. TikTok is present in the connector layer as an additional supported source."],
            ["Primary output", "Email report and Databricks dashboard-ready tables. Future login dashboards can sit on the same data foundation."],
            ["Business model", "Managed revenue intelligence package: attribution execution, source QA, revenue reporting, monthly recommendations, and agency portfolio views."],
        ],
        [2100, 7260],
    )
    add_heading(doc, "Attribution Models", 2)
    add_table(
        doc,
        ["Model", "How credit is assigned", "Best fit"],
        [
            ["last_touch", "All revenue credit goes to the final eligible touchpoint before conversion.", "Shorter cycles, direct response, retargeting, or when the client wants to optimize closing motions."],
            ["first_touch", "All credit goes to the first eligible touchpoint in the journey.", "Demand creation, SEO/content/social awareness, and long cycles where first source matters."],
            ["linear", "Credit is split equally across all eligible touchpoints.", "Balanced multi-touch cycles where the client wants a neutral view."],
            ["time_decay", "Recent touches receive more weight than older touches.", "Longer cycles with active nurture where late-stage momentum matters."],
            ["u_shape", "Heavier credit goes to first and last touch; middle touches share the rest.", "Lead generation journeys where both source creation and conversion touch matter."],
            ["w_shape", "Heavier credit goes to first touch, lead creation, and opportunity/close touch.", "B2B sales cycles with meaningful lead and opportunity stages."],
        ],
        [1500, 4200, 3660],
    )


def section_architecture(doc: Document) -> None:
    add_heading(doc, "3. System Architecture", 1)
    add_para(
        doc,
        "The production path intentionally separates compute from durable analytics. Cloud Run performs external API calls and orchestration because Databricks Serverless egress constraints can block DNS/API access to ad and payment systems. Databricks remains the durable SQL/Delta data layer.",
    )
    add_table(
        doc,
        ["Layer", "Component", "Files or service"],
        [
            ["Operator UI", "Streamlit Command Center", r"attribution_agent\attribution_agent\app.py"],
            ["Pipeline compute", "Cloud Run Job attribution-pipeline", r"Dockerfile, deploy.sh, scripts\arie_deploy_cloud_run.ps1"],
            ["UI hosting", "Cloud Run Service attribution-ui", r"Streamlit running in the same container image"],
            ["Scheduling", "Cloud Scheduler attribution-monthly", "Monthly trigger, default 9am ET on the 1st"],
            ["Data layer", "Databricks SQL/Delta", r"utils\databricks_writer.py"],
            ["Client registry", "GCS-mounted clients.json or Databricks ops client_registry", r"config\client_config.py"],
            ["Secrets", "GCP Secret Manager plus local .env fallback", r"utils\secrets.py"],
            ["Reporting AI", "Claude via ModelGateway", r"agents\insight\insight_agent.py, utils\model_gateway.py"],
            ["Communications", "Email and Telegram", r"agents\comms\comms_agent.py, agents\control\arie_bot.py"],
            ["External interfaces", "FastAPI, A2A, MCP", "api, agents\\a2a, mcp_server"],
        ],
        [1600, 3100, 4660],
    )
    add_heading(doc, "Run Flow", 2)
    for item in [
        "A trigger starts a run: Command Center, FastAPI, MCP, Telegram/CLI workflow, Cloud Scheduler, or direct Cloud Run Job execution.",
        "ARIE loads agency/client configuration and resolves credentials from Secret Manager or local environment fallbacks.",
        "The ingest flow pulls enabled sources, validates rows, normalizes ad touchpoints, matches revenue journeys, applies the selected attribution model, and writes results.",
        "The agency flow refreshes attribution SQL, generates the insight report, optionally runs governance review, records run history, and sends/skips email depending on dry-run mode.",
        "Observability tables track run state, row counts, warnings, errors, reports, costs, approvals, audit events, Telegram triggers, and model gateway activity.",
    ]:
        add_number(doc, item)


def section_pipeline(doc: Document) -> None:
    add_heading(doc, "4. Attribution Pipeline", 1)
    add_table(
        doc,
        ["Stage", "What happens", "Important files"],
        [
            ["Setup", "Load client config, resolve tokens, ensure schema/tables, initialize checkpoints.", r"flows\ingest_flow.py, utils\checkpoint.py"],
            ["Ingest", "Pull ad, CRM, and payment sources with retries. Failed sources are recorded as source_failures instead of killing the entire run.", r"agents\ingest\*_connector.py"],
            ["Validate", "Check source-level issues including missing spend, schema drift, spend drops, zero spend, and partial data.", r"agents\ingest\validator.py"],
            ["Normalize", "Convert platform-specific ad records into a shared ad/touchpoint model.", r"agents\ingest\ad_sources.py"],
            ["Attribute", "Build deal/revenue journeys and assign revenue credit using the selected model.", r"attribution_engine.py, attribution_models.py"],
            ["Persist", "Write raw/normalized/attributed tables, run records, warnings, and outputs to Databricks.", r"utils\databricks_writer.py"],
            ["Report", "Create client-facing insight report from refreshed attribution output and send email unless dry-run is enabled.", "agents\\insight, agents\\comms"],
        ],
        [1500, 5200, 2660],
    )
    add_heading(doc, "Validation And Reliability Behavior", 2)
    for item in [
        "Enabled sources with missing/expired credentials degrade to partial runs and are written into source_failures.",
        "Pipeline steps are checkpointed so a resume run can skip completed work.",
        "PipelineSaga provides a rollback pattern around per-client execution failures.",
        "Dry-run mode performs attribution/report preview while skipping email delivery.",
        "Governance review is advisory: warnings are logged but do not block report generation.",
    ]:
        add_bullet(doc, item)


def section_data_layer(doc: Document) -> None:
    add_heading(doc, "5. Data Layer And Tables", 1)
    add_para(
        doc,
        "Databricks is the system of record for source data, attribution outputs, operational telemetry, and report-ready tables. Client business data is written to client schemas; shared operational state lives in the ops schema, defaulting to workspace.attribution_ops.",
    )
    add_table(
        doc,
        ["Concern", "Tables or outputs"],
        [
            ["Run tracking", "pipeline_runs, pipeline_checkpoints, idempotency_store"],
            ["Reporting", "insight_reports, approval_queue"],
            ["Governance and audit", "audit_log"],
            ["Cost and caching", "cost_ledger, semantic_cache"],
            ["Agent context", "agent_memory"],
            ["Evaluations", "eval_golden_dataset, eval_results"],
            ["Experiments", "ab_experiments, ab_assignments"],
            ["Clients", "client_registry"],
            ["Telegram operations", "telegram_events"],
            ["Client attribution outputs", "normalized ad data, attribution results, closed_revenue_attribution, agency dashboard and benchmark transforms"],
        ],
        [2400, 6960],
    )
    add_heading(doc, "Important SQL Transforms", 2)
    add_table(
        doc,
        ["Transform", "Purpose"],
        [
            [r"transforms\closed_revenue_attribution.sql", "Applies selected model logic to closed-won CRM revenue and touchpoint/deal journeys."],
            [r"transforms\attribution_model.sql", "Builds attribution model outputs for reporting."],
            [r"transforms\attribution_model_with_payments.sql", "Adds Stripe payment/revenue enrichment where available."],
            [r"transforms\agency_dashboard.sql", "Maintains agency-level dashboard tables."],
            [r"transforms\agency_benchmark.sql", "Creates cross-client agency benchmarks."],
        ],
        [3600, 5760],
    )


def section_command_center(doc: Document) -> None:
    add_heading(doc, "6. Command Center And Client Setup", 1)
    add_para(
        doc,
        "The ARIE Command Center is the internal operator surface. It is currently a Streamlit application with tabs for execution, client onboarding, outreach, and observability. It is sufficient for internal operation today and can later be replaced or wrapped by a login-based agency/customer portal.",
    )
    add_table(
        doc,
        ["Tab", "Operator purpose"],
        [
            ["Pipeline", "Select agency/client scope, dry-run/live mode, attribution model, run mode, and submit the Cloud Run pipeline command."],
            ["Clients", "Add, remove, and update clients; capture business details, account IDs, attribution defaults, report emails, and credentials."],
            ["Outreach", "Generate and manage operator outreach/prospecting actions associated with ARIE/N8iV workflows."],
            ["Observability", "Inspect recent pipeline health, row counts, warnings, costs, evals, audit activity, and Telegram triggers."],
        ],
        [1700, 7660],
    )
    add_heading(doc, "Client Onboarding Data Captured", 2)
    add_table(
        doc,
        ["Category", "Fields"],
        [
            ["Identity", "client_name, client_display_name, agency_id, Databricks schema"],
            ["Run defaults", "default attribution_model, lookback_days, report email"],
            ["Meta", "enabled flag, ad account ID, access token secret"],
            ["Google Ads", "enabled flag, customer ID, refresh token secret plus global developer/client credentials"],
            ["LinkedIn", "enabled flag, ad account ID, access token secret"],
            ["TikTok", "enabled flag, advertiser ID, access token secret"],
            ["HubSpot", "enabled flag, pipeline ID, access token secret"],
            ["Stripe", "enabled flag, account ID, secret key secret"],
        ],
        [2200, 7160],
    )
    add_callout(
        doc,
        "Credential storage rule",
        "Credential values entered in the UI/API are written to GCP Secret Manager. The client registry stores only Secret Manager IDs, using names like attr-prod-<client-id>-<credential-name> unless the prefix/environment is overridden.",
    )


def section_interfaces(doc: Document) -> None:
    add_heading(doc, "7. API, A2A, MCP, And Telegram Interfaces", 1)
    add_heading(doc, "FastAPI", 2)
    add_para(
        doc,
        "FastAPI exposes administrative and reporting operations. Auth uses the X-API-Key header. API_KEY_ADMIN maps to admin; per-client API_KEY_<CLIENT_ID> keys map to analyst-level access.",
    )
    add_table(
        doc,
        ["Route", "Purpose"],
        [
            ["GET /health", "Service health and build metadata."],
            ["POST /api/v1/pipeline/run", "Submit Cloud Run or Databricks pipeline execution with agency_id, client_ids, dry_run, attribution_model, and run_mode."],
            ["GET /api/v1/pipeline/runs", "List recent pipeline runs from Databricks ops history."],
            ["GET /api/v1/pipeline/runs/{run_id}", "Fetch a specific run record."],
            ["GET /api/v1/clients", "List clients."],
            ["GET /api/v1/clients/{client_id}", "Fetch a client config summary."],
            ["POST /api/v1/clients", "Create a client and store credential values in Secret Manager."],
            ["PUT /api/v1/clients/{client_id}", "Update a client while preserving blank credential fields."],
            ["DELETE /api/v1/clients/{client_id}", "Delete custom clients; built-in clients are protected."],
            ["GET /api/v1/reports/{client_id}/latest", "Fetch latest stored report."],
            ["GET /api/v1/reports/{client_id}", "List report history."],
            ["POST /api/v1/reports/{client_id}/generate", "Start a background report generation task."],
            ["GET /api/v1/approvals", "List pending approval actions."],
            ["POST /api/v1/approvals/{action_id}/resolve", "Approve or reject a pending action."],
        ],
        [3300, 6060],
    )
    add_heading(doc, "A2A And MCP", 2)
    add_table(
        doc,
        ["Interface", "Current purpose"],
        [
            ["A2A", "Agent discovery at /.well-known/agent-cards and dispatch at /a2a/dispatch for data-quality, revenue-analyst, executive-reporting, and governance-reviewer style calls."],
            ["MCP", "FastMCP stdio server with tools for listing clients, reading client config, checking pipeline status, generating reports, running ingest flow, cost summary, and recent memories."],
        ],
        [1600, 7760],
    )
    add_heading(doc, "Telegram Bot", 2)
    add_para(
        doc,
        "The ARIE Telegram bot supports status, overview, prospect/outreach, and pipeline commands. Local changes add durable trigger logging and a cross-process lock because Telegram allows only one getUpdates consumer per bot token.",
    )
    for item in [
        "Start standalone listener through arie_start.ps1 or scripts/arie_open_command_center.ps1 -StartBot.",
        "Do not start both Streamlit-embedded bot and standalone bot at the same time.",
        "The app defaults to not auto-starting the bot; set ARIE_COMMAND_CENTER_START_BOT=true only when intentionally embedding it.",
        "Recent bot triggers can be reviewed in the Command Center Observability tab after Databricks logging succeeds.",
    ]:
        add_bullet(doc, item)


def section_integrations(doc: Document) -> None:
    add_heading(doc, "8. Integrations And Credentials", 1)
    add_table(
        doc,
        ["Source", "Role in attribution", "Credential handling"],
        [
            ["Meta", "Ad spend/touchpoints from Meta Graph API.", "Per-client access token in Secret Manager."],
            ["Google Ads", "Ad spend/touchpoints from Google Ads.", "Global developer/client credentials plus per-client refresh token."],
            ["LinkedIn Ads", "Ad spend/touchpoints from LinkedIn.", "Per-client access token."],
            ["TikTok Ads", "Optional additional ad source present in code.", "Per-client access token."],
            ["HubSpot", "Primary CRM source of truth for B2B closed-won revenue.", "Per-client access token and pipeline ID."],
            ["Stripe", "Collected revenue/refund enrichment and true ROI support.", "Per-client secret key/account ID."],
            ["Claude", "Report generation, data quality, governance review, and agent reasoning.", "Global ANTHROPIC_API_KEY."],
            ["Gmail/SendGrid", "Client-facing email report delivery.", "GMAIL_SENDER and GMAIL_APP_PASSWORD; SENDGRID_API_KEY is optional."],
            ["Databricks", "Durable SQL/Delta system of record.", "Global DATABRICKS_SERVER_HOSTNAME, HTTP_PATH, TOKEN."],
            ["GCP", "Cloud Run, Scheduler, Secret Manager, Artifact Registry, Storage.", "gcloud operator auth plus service account/IAM setup."],
        ],
        [1600, 4100, 3660],
    )
    add_heading(doc, "Global Secrets Mentioned By The Build", 2)
    for item in [
        "DATABRICKS_SERVER_HOSTNAME, DATABRICKS_HTTP_PATH, DATABRICKS_TOKEN",
        "ANTHROPIC_API_KEY",
        "API_KEY_ADMIN",
        "GMAIL_SENDER, GMAIL_APP_PASSWORD",
        "SENDGRID_API_KEY, optional",
        "GOOGLE_ADS_DEVELOPER_TOKEN, GOOGLE_ADS_CLIENT_ID, GOOGLE_ADS_CLIENT_SECRET, GOOGLE_ADS_LOGIN_CUSTOMER_ID",
    ]:
        add_bullet(doc, item)
    add_callout(
        doc,
        "Security note",
        "This document intentionally lists secret names and storage patterns only. It does not include secret values, admin keys, tokens, API credentials, or client credentials.",
        fill=PAPER,
    )


def section_deployment(doc: Document) -> None:
    add_heading(doc, "9. Deployment And Operations", 1)
    add_table(
        doc,
        ["Operation", "Command or path"],
        [
            ["Preflight", r".\scripts\arie_pilot_preflight.ps1 -ProjectId <project> -AfterDeploy"],
            ["Seed secrets", r".\scripts\arie_seed_secrets.ps1 -ProjectId <project>"],
            ["Deploy Cloud Run", r".\scripts\arie_deploy_cloud_run.ps1 -ProjectId <project> -OperatorPrincipal user:you@example.com"],
            ["Open private UI", r"gcloud run services proxy attribution-ui --project <project> --region us-central1 --port 8080"],
            ["Cloud Run dry run", r".\scripts\arie_cloud_run_dry_run.ps1 -ProjectId <project> -AgencyId <agency_id> -AttributionModel w_shape -ClientIds <client_id>"],
            ["Local Streamlit", r".\arie_start.ps1 or .\scripts\arie_open_command_center.ps1 -Mode Local -StartBot"],
            ["Stop local services", r".\arie_stop.ps1"],
            ["Desktop shortcut", r".\scripts\install_arie_desktop_shortcut.ps1 -Mode Local -StartBot"],
        ],
        [2100, 7260],
    )
    add_heading(doc, "Cloud Run Production Path", 2)
    for item in [
        "One image is built from the root Dockerfile and pushed to Artifact Registry.",
        "Cloud Run service attribution-ui hosts the Streamlit Command Center.",
        "Cloud Run job attribution-pipeline runs the agency/client pipeline.",
        "Cloud Scheduler attribution-monthly runs the job on a schedule.",
        "GCS bucket <project>-attribution-registry stores clients.json when using local registry backend in Cloud Run.",
        "Secret Manager injects global environment secrets at deploy time and stores per-client source credentials at onboarding time.",
    ]:
        add_bullet(doc, item)


def section_tests_issues(doc: Document) -> None:
    add_heading(doc, "10. Tests, Usability, And Known Issues", 1)
    add_heading(doc, "Automated Coverage Present", 2)
    add_table(
        doc,
        ["Test area", "Files"],
        [
            ["Attribution allocation", r"tests\test_attribution_models.py, tests\test_attribution_engine.py"],
            ["Ad source normalization", r"tests\test_ad_sources.py"],
            ["Connector contracts", r"tests\test_connector_contracts.py"],
            ["Client registry", r"tests\test_client_registry_store.py"],
            ["Cloud Run args", r"tests\test_cloud_run_job.py"],
            ["A2A dispatch", r"tests\test_a2a_dispatch.py"],
            ["Secret redaction and PII", r"tests\test_secret_redaction.py, tests\test_pii_masker.py"],
            ["Model gateway", r"tests\test_model_gateway.py"],
            ["Demo flows", r"tests\test_demo_engine.py"],
        ],
        [2400, 6960],
    )
    add_heading(doc, "Recorded Usability Result", 2)
    add_callout(
        doc,
        "Latest UI smoke test",
        "reports/usability/2026-07-15T02-49-25-712Z recorded 38 passing checks, 2 warnings, and 0 failures. The warnings were browser console Vega-Lite version/fit-x messages. The test covered desktop and mobile loading, Pipeline, Clients, Outreach, Observability, recent runs interaction, and form filling without saving.",
        fill=PAPER,
    )
    add_heading(doc, "Known Issues And Risks", 2)
    add_table(
        doc,
        ["Risk", "Impact", "Recommended action"],
        [
            ["Local runtime stale", "PID files exist but local UI/API were not reachable during documentation.", "Run arie_stop.ps1, then arie_start.ps1 or desktop shortcut; verify 8501 and 8081 health."],
            ["Approvals route bug", "POST /api/v1/approvals/{action_id}/resolve likely raises TypeError because actor= is passed to a helper expecting resolved_by.", "Patch call to resolved_by=\"api\" and add route test."],
            ["GCP auth/tooling", "Operator PowerShell is authenticated for GCP; the Codex sandbox may still lack gcloud on PATH.", "Run deployed checks from the authenticated terminal and keep auth refreshed before Cloud Run or Secret Manager work."],
            ["Live pipeline not revalidated in this pass", "No live Cloud Run job, email, or source API execution was triggered.", "Run one N8iV dry run, inspect Databricks outputs, then run a controlled live send."],
            ["Website/DNS ambiguity", "Public website may still resolve to old/Wix/placeholder content.", "Verify Vercel deployment domains, DNS records, Google Search Console, sitemap, and indexing."],
            ["GCS client registry concurrency", "GCS FUSE clients.json has no multi-writer protection.", "Fine for one operator now; use Databricks registry or a transactional store before agency logins."],
            ["Secrets rotation behavior", "Env-injected Cloud Run UI secrets need redeploy/recycle to pick up latest versions.", "Document rotation SOP and prefer runtime Secret Manager reads for per-client secrets."],
        ],
        [2300, 3000, 4060],
    )


def section_checklist(doc: Document) -> None:
    add_heading(doc, "11. Production Readiness Checklist", 1)
    add_table(
        doc,
        ["Priority", "Item", "Done when"],
        [
            ["P0", "Patch approvals route and add regression test.", "Approval resolve endpoint returns success/failure correctly and logs audit event."],
            ["P0", "Restart and verify local ARIE app/API.", "Command Center loads, FastAPI /health returns ok, stale PID files cleaned."],
            ["P0", "Reauthenticate gcloud and run deployed preflight.", "Cloud Run service/job, Scheduler, bucket, and required secrets pass."],
            ["P0", "Run N8iV dry-run pilot.", "pipeline_runs shows success/partial with selected model, row counts, warnings, and email_sent=false."],
            ["P0", "Verify Databricks revenue tables.", "Raw, normalized, closed_revenue_attribution, and report output tables refresh before report generation."],
            ["P0", "Send one approved internal live report.", "Email delivery succeeds to operator/test inbox and content matches selected model."],
            ["P1", "Complete client onboarding SOP.", "All required fields, source credential instructions, consent/access steps, and failure handling are documented."],
            ["P1", "Add source-level alert thresholds.", "Missing UTMs, missing emails, missing closed-won amounts, unmatched payments, and high unattributed revenue are visible in UI/report."],
            ["P1", "Lock down Cloud Run IAM.", "Private UI access works through proxy or authenticated browser; no public secret exposure."],
            ["P1", "Clean repo artifacts.", "Duplicate old folders/prototypes are quarantined or intentionally documented."],
            ["P2", "Prepare agency-login future path.", "Choose auth/tenant model, registry persistence strategy, and user roles before external access."],
        ],
        [900, 3600, 4860],
    )
    add_callout(
        doc,
        "Fastest path to sell pilots",
        "Keep the Command Center internal, onboard fewer than 10 pilot accounts, use HubSpot closed-won revenue as the anchor, enrich with Stripe only where available, run dry-run first, then deliver monthly email reports with clear model-specific language and operator QA notes.",
    )


def section_commands(doc: Document) -> None:
    add_heading(doc, "12. Useful Commands", 1)
    commands = [
        ("Local app", r".\arie_start.ps1"),
        ("Stop local app", r".\arie_stop.ps1"),
        ("Open desktop-style app", r".\scripts\arie_open_command_center.ps1 -Mode Local -StartBot"),
        ("Install shortcut", r".\scripts\install_arie_desktop_shortcut.ps1 -Mode Local -StartBot"),
        ("Preflight", r".\scripts\arie_pilot_preflight.ps1 -ProjectId n8iv-analytics-production -AfterDeploy"),
        ("Seed secrets", r".\scripts\arie_seed_secrets.ps1 -ProjectId n8iv-analytics-production"),
        ("Deploy", r".\scripts\arie_deploy_cloud_run.ps1 -ProjectId n8iv-analytics-production -OperatorPrincipal user:zajen@n8ivpromotions.com"),
        ("Cloud proxy", r"gcloud run services proxy attribution-ui --project n8iv-analytics-production --region us-central1 --port 8080"),
        ("Dry run", r".\scripts\arie_cloud_run_dry_run.ps1 -ProjectId n8iv-analytics-production -AgencyId <agency_id> -AttributionModel w_shape -ClientIds <client_id>"),
        ("FastAPI docs", r"http://127.0.0.1:8081/docs"),
        ("Streamlit local", r"http://127.0.0.1:8501"),
    ]
    add_table(doc, ["Task", "Command"], [[a, b] for a, b in commands], [2100, 7260])
    add_para(
        doc,
        "End of build documentation. Treat this document as the handoff snapshot for the current ARIE repository state, not as a replacement for source-controlled README/RUNBOOK updates.",
        size=10.5,
        color=MUTED,
        after=0,
    )


def build_doc() -> Path:
    doc = Document()
    configure_doc(doc)
    doc.core_properties.title = "ARIE Build Documentation"
    doc.core_properties.subject = "Automatic Revenue Intelligence Engine build documentation"
    doc.core_properties.author = "N8iV Promotions"
    doc.core_properties.comments = "Generated documentation for the attribution-agent build. No secrets included."

    add_cover(doc)
    add_document_map(doc)
    section_current_snapshot(doc)
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    section_product(doc)
    section_architecture(doc)
    section_pipeline(doc)
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    section_data_layer(doc)
    section_command_center(doc)
    section_interfaces(doc)
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    section_integrations(doc)
    section_deployment(doc)
    section_tests_issues(doc)
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    section_checklist(doc)
    section_commands(doc)
    set_running_footer(doc)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    print(build_doc())
